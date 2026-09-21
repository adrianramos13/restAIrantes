"""
Extrae reseñas y localización de restaurantes usando la API de Outscraper.

USO:
    1. Instala dependencias:
         pip install outscraper openpyxl pandas

    2. Define tu API key como variable de entorno (no la pegues en el código):
         En Mac/Linux:   export OUTSCRAPER_API_KEY="tu_api_key_aqui"
         En Windows:     set OUTSCRAPER_API_KEY=tu_api_key_aqui

    3. Ajusta las constantes de configuración más abajo si hace falta
       (ruta del excel de entrada, ciudad, número de reseñas por restaurante...).

    4. Ejecuta:
         python extraer_resenas.py
"""

import os
import sys
import time
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from outscraper import OutscraperClient


# ----------------------- CONFIGURACIÓN ----------------------- #

EXCELS_DIR = Path(__file__).resolve().parent.parent / "excels"

INPUT_XLSX = EXCELS_DIR / "restaurantes_novia.xlsx"      # Excel con la columna "Nombre"
OUTPUT_XLSX = EXCELS_DIR / "restaurantes_con_resenas.xlsx"

CIUDAD = "Madrid, España"   # se añade a cada nombre para desambiguar la búsqueda
REVIEWS_LIMIT = 30          # reseñas por restaurante (entre 20 y 50 según lo hablado)
IDIOMA = "es"
REGION = "ES"
BATCH_SIZE = 25             # nº de restaurantes por petición (evita el error 414 URI Too Long)

# --------------------------------------------------------------- #


def cargar_nombres(ruta_excel):
    """Lee la columna 'Nombre' del excel generado en el paso anterior."""
    df = pd.read_excel(ruta_excel)
    if "Nombre" not in df.columns:
        raise ValueError(f"No encuentro una columna 'Nombre' en {ruta_excel}")
    nombres = df["Nombre"].dropna().astype(str).tolist()
    return nombres


def construir_queries(nombres, ciudad):
    """Combina cada nombre con la ciudad para que Outscraper no confunda
    restaurantes homónimos de otras ciudades."""
    return [f"{nombre}, {ciudad}" for nombre in nombres]


def pedir_datos_outscraper(client, queries, batch_size=BATCH_SIZE):
    """
    Llama a la API de reseñas de Outscraper EN LOTES.

    El endpoint usa peticiones GET con todas las queries metidas en la URL,
    así que mandarlas todas juntas (226 nombres largos) supera el límite de
    longitud de URL del servidor (error 414). Troceamos en lotes pequeños.

    limit=1 -> nos quedamos con el primer resultado (más relevante) por query,
    ya que cada query ya identifica un restaurante concreto.
    """
    print(f"Consultando Outscraper para {len(queries)} restaurantes, "
          f"en lotes de {batch_size} (esto puede tardar varios minutos)...")

    todos_resultados = []
    total_lotes = (len(queries) + batch_size - 1) // batch_size

    for i in range(0, len(queries), batch_size):
        lote = queries[i:i + batch_size]
        num_lote = i // batch_size + 1
        print(f"  Lote {num_lote}/{total_lotes} ({len(lote)} restaurantes)...")

        try:
            resultados_lote = client.google_maps_reviews(
                lote,
                reviews_limit=REVIEWS_LIMIT,
                limit=1,
                language=IDIOMA,
                region=REGION,
            )
            todos_resultados.extend(resultados_lote or [])
        except Exception as e:
            print(f"    [ERROR] Falló el lote {num_lote}: {e}")
            # Rellenamos con None para no descuadrar el emparejamiento
            # posicional con nombres_originales en procesar_resultados().
            todos_resultados.extend([None] * len(lote))

        # Pequeña pausa entre lotes para no saturar la API.
        if i + batch_size < len(queries):
            time.sleep(1)

    return todos_resultados


def procesar_resultados(nombres_originales, resultados):
    """
    Separa los resultados en dos listas de filas:
      - restaurantes: una fila por local (info general + localización)
      - resenas: una fila por reseña individual

    Como google_maps_reviews puede devolver menos elementos que queries
    (si algún restaurante no se encontró), emparejamos por posición y
    avisamos de los que falten.
    """
    filas_restaurantes = []
    filas_resenas = []
    no_encontrados = []

    # Outscraper no siempre devuelve exactamente 1 resultado por query si
    # no encuentra nada, así que vamos consumiendo la lista en orden y
    # comparamos cuántos "huecos" quedan al final.
    resultados = resultados or []

    if len(resultados) < len(nombres_originales):
        print(f"[AVISO] Se esperaban {len(nombres_originales)} resultados "
              f"y llegaron {len(resultados)}. Algunos restaurantes no se "
              f"habrán podido localizar; revisa 'no_encontrados' al final.")

    for i, nombre_original in enumerate(nombres_originales):
        if i >= len(resultados) or not resultados[i]:
            no_encontrados.append(nombre_original)
            continue

        lugar = resultados[i]

        filas_restaurantes.append({
            "Nombre (excel original)": nombre_original,
            "Nombre (Google Maps)": lugar.get("name", ""),
            "Dirección": lugar.get("full_address", ""),
            "Latitud": lugar.get("latitude", ""),
            "Longitud": lugar.get("longitude", ""),
            "Rating global": lugar.get("rating", ""),
            "Nº reseñas totales (Google)": lugar.get("reviews", ""),
            "Teléfono": lugar.get("phone", ""),
            "Web": lugar.get("site", ""),
        })

        for review in lugar.get("reviews_data", []) or []:
            filas_resenas.append({
                "Restaurante": nombre_original,
                "Autor": review.get("autor_name", ""),
                "Puntuación": review.get("review_rating", ""),
                "Fecha": review.get("review_datetime_utc", ""),
                "Texto": review.get("review_text", ""),
                "Respuesta del propietario": review.get("owner_answer", ""),
            })

    return filas_restaurantes, filas_resenas, no_encontrados


def guardar_excel(filas_restaurantes, filas_resenas, no_encontrados, ruta_salida):
    wb = openpyxl.Workbook()

    # --- Hoja 1: Restaurantes ---
    ws1 = wb.active
    ws1.title = "Restaurantes"
    if filas_restaurantes:
        headers1 = list(filas_restaurantes[0].keys())
        ws1.append(headers1)
        for fila in filas_restaurantes:
            ws1.append([fila[h] for h in headers1])

    # --- Hoja 2: Reseñas ---
    ws2 = wb.create_sheet("Reseñas")
    if filas_resenas:
        headers2 = list(filas_resenas[0].keys())
        ws2.append(headers2)
        for fila in filas_resenas:
            ws2.append([fila[h] for h in headers2])

    # --- Hoja 3: No encontrados (si aplica) ---
    if no_encontrados:
        ws3 = wb.create_sheet("No encontrados")
        ws3.append(["Nombre"])
        for nombre in no_encontrados:
            ws3.append([nombre])

    # Estilo básico: cabecera en negrita + fuente Arial en todo el libro
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial")
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions

    # Anchos de columna razonables
    for ws, widths in [
        (ws1, {"A": 30, "B": 30, "C": 40, "D": 12, "E": 12, "F": 12, "G": 15, "H": 18, "I": 30}),
        (ws2, {"A": 30, "B": 20, "C": 12, "D": 18, "E": 60, "F": 40}),
    ]:
        for col, w in widths.items():
            ws.column_dimensions[col].width = w

    wb.save(ruta_salida)
    print(f"\nGuardado: {ruta_salida}")
    print(f"  - {len(filas_restaurantes)} restaurantes")
    print(f"  - {len(filas_resenas)} reseñas")
    if no_encontrados:
        print(f"  - {len(no_encontrados)} restaurantes NO encontrados (ver hoja 'No encontrados')")


def main():
    api_key = os.environ.get("OUTSCRAPER_API_KEY")
    if not api_key:
        sys.exit(
            "ERROR: No encuentro la variable de entorno OUTSCRAPER_API_KEY.\n"
            "Defínela antes de ejecutar el script, por ejemplo:\n"
            '  export OUTSCRAPER_API_KEY="tu_api_key_aqui"   (Mac/Linux)\n'
            "  set OUTSCRAPER_API_KEY=tu_api_key_aqui        (Windows)"
        )

    if not os.path.exists(INPUT_XLSX):
        sys.exit(f"ERROR: No encuentro el archivo de entrada '{INPUT_XLSX}'.")

    nombres = cargar_nombres(INPUT_XLSX)
    print(f"Leídos {len(nombres)} restaurantes de '{INPUT_XLSX}'.")

    queries = construir_queries(nombres, CIUDAD)

    client = OutscraperClient(api_key=api_key)

    t0 = time.time()
    resultados = pedir_datos_outscraper(client, queries)
    print(f"Consulta completada en {time.time() - t0:.1f}s.")

    filas_restaurantes, filas_resenas, no_encontrados = procesar_resultados(nombres, resultados)

    guardar_excel(filas_restaurantes, filas_resenas, no_encontrados, OUTPUT_XLSX)


if __name__ == "__main__":
    main()