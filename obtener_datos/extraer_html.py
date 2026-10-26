"""
Extrae los datos de restaurantes desde el HTML de tu lista de Google Maps
(guardado en la carpeta origen/) y genera excels/restaurantes_v1.xlsx.

CAMPOS EXTRAÍDOS:
    - Nombre
    - Puntuación
    - Nº Reseñas
    - Rango de precios
    - Tipo de cocina
    - Estado           (vacío normalmente; "Cerrado temporalmente" o
                         "Cerrado permanentemente" si Google Maps lo marca así)
    - Imagen URL / Imagen archivo
                       las URLs de fotos de Google caducan en unas semanas, así que
                       cada foto se descarga UNA vez a imagenes/<nombre>.jpg (que se
                       sube al repo) y la app usa ese archivo. Los restaurantes que ya
                       tienen foto descargada no se vuelven a descargar: al añadir
                       restaurantes nuevos, solo se bajan las fotos de los nuevos.
                       Si la URL del HTML ya ha caducado, se busca el restaurante en
                       Google Maps con Playwright (sin coste) para sacar una URL nueva.

CÓMO OBTENER EL HTML:
    Automático (recomendado): comparte tu lista en Google Maps (Compartir -> copiar
    enlace) y ejecuta  python extraer_html.py --url "<enlace>". Abre la lista con
    Playwright, hace scroll hasta cargarla entera y la guarda en origen/restaurantes.html.
    Si la página no trae restaurantes, no toca ese archivo: guarda lo descargado en
    origen/restaurantes_descarga_fallida.html y se para.

    A mano:
    1. Abre tu lista de Google Maps en el navegador.
    2. Haz scroll hasta cargar TODOS los restaurantes (Google Maps carga
       la lista poco a poco al hacer scroll).
    3. Clic derecho -> "Inspeccionar" -> en el árbol del DOM busca el
       contenedor de la lista y copia su HTML ("Copy outerHTML"), o guarda
       la página completa con Ctrl+S.
    4. Guarda ese HTML como origen/restaurantes.html

USO:
    python extraer_html.py --url "https://maps.app.goo.gl/..."   (descarga la lista)
    python extraer_html.py                                       (usa origen/restaurantes.html)
"""

import asyncio
import re
import unicodedata
import urllib.parse
from pathlib import Path

import requests

from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill


# ----------------------- CONFIGURACIÓN ----------------------- #

BASE_DIR = Path(__file__).resolve().parent
INPUT_HTML = BASE_DIR.parent / "origen" / "restaurantes.html"
OUTPUT_XLSX = BASE_DIR.parent / "excels" / "restaurantes_v1.xlsx"
IMAGENES_DIR = BASE_DIR.parent / "imagenes"
TAMANO_IMAGEN = "=w450-h450"   # la miniatura del HTML es diminuta (=w80-h142); se pide más grande

# Descarga automática de la lista (--url)
HTML_DESCARGA_FALLIDA = BASE_DIR.parent / "origen" / "restaurantes_descarga_fallida.html"
SELECTOR_ENTRADA_LISTA = "button.SMP2wb"   # cada restaurante de la lista: el mismo que lee extraer_restaurantes()
RONDAS_SIN_NUEVOS = 5                     # se deja de hacer scroll tras 5 rondas seguidas sin restaurantes nuevos
MAX_RONDAS_SCROLL = 300                   # tope por si la lista nunca deja de "cargar"
PAUSA_POR_ENTRADA_MS = 150                # al recorrer la lista, para que Google rellene cada entrada
# Sesión de Google guardada (solo en este ordenador, fuera de git): con ella la lista trae el precio.
PERFIL_CHROME_DIR = BASE_DIR.parent / "origen" / "perfil_chrome"
ESPERA_INICIO_SESION_S = 300
SIN_VENTANA = False   # Chrome sin ventana (procesar_solicitudes.py --sin-ventana, la tarea de Windows)

# --------------------------------------------------------------- #


def clean(text):
    return re.sub(r"\s+", " ", text).strip()


def nombre_archivo_imagen(nombre):
    """'K'era | Restaurante Georgiano' -> 'kera-restaurante-georgiano.jpg'. Depende solo del
    nombre (no del ID, que cambia al añadir restaurantes), así la foto sigue siendo la suya."""
    sin_tildes = unicodedata.normalize("NFKD", nombre).encode("ascii", "ignore").decode().lower()
    slug = re.sub(r"[^a-z0-9]+", "-", sin_tildes.replace("'", "")).strip("-")[:80]
    return f"{slug or 'sin-nombre'}.jpg"


def bajar_imagen(url, ruta):
    """Descarga `url` (pidiendo TAMANO_IMAGEN) a `ruta`. Lanza excepción si no es una imagen."""
    r = requests.get(re.sub(r"=w\d+-h\d+", TAMANO_IMAGEN, url), timeout=15)
    r.raise_for_status()
    if not r.headers.get("content-type", "").startswith("image/"):
        raise ValueError(f"no es una imagen ({r.headers.get('content-type')})")
    ruta.write_bytes(r.content)


async def _urls_frescas(nombres):
    """Abre la ficha de cada restaurante en Google Maps (con Playwright, como extraer_resenas.py)
    y devuelve {nombre: URL de su foto principal}, recién generada y por tanto aún válida."""
    from playwright.async_api import async_playwright
    from extraer_resenas import aceptar_cookies

    urls = {}
    async with async_playwright() as p:
        navegador = await p.chromium.launch(headless=True)
        page = await navegador.new_page(locale="es-ES")
        for i, nombre in enumerate(nombres, start=1):
            print(f"  [{i}/{len(nombres)}] {nombre}")
            consulta = urllib.parse.quote(f"{nombre}, Madrid, España")
            try:
                await page.goto(f"https://www.google.com/maps/search/?api=1&query={consulta}",
                                wait_until="domcontentloaded")
                await aceptar_cookies(page)
                foto = page.locator('button[jsaction*="heroHeaderImage"] img').first
                primer_resultado = page.locator("a.hfpxzc").first
                # La búsqueda abre directamente la ficha o, si hay varias coincidencias, una lista.
                await foto.or_(primer_resultado).wait_for(timeout=15000)
                if not await foto.count():
                    await primer_resultado.click()
                    await foto.wait_for(timeout=15000)
                urls[nombre] = await foto.get_attribute("src")
            except Exception as e:
                print(f"      sin foto: {type(e).__name__}")
        await navegador.close()
    return urls


def descargar_imagenes(filas):
    """Descarga a imagenes/ las fotos que aún no estén, y rellena 'Imagen archivo' en cada fila
    (vacío si no hay foto ni se ha podido bajar; la app entonces prueba con la URL).
    Si la URL del HTML ha caducado (HTML antiguo), busca el restaurante en Google Maps para
    sacar una URL nueva de su foto."""
    IMAGENES_DIR.mkdir(parents=True, exist_ok=True)
    nuevas, pendientes = 0, []
    for fila in filas:
        ruta = IMAGENES_DIR / nombre_archivo_imagen(fila["Nombre"])
        if not ruta.exists():
            try:
                bajar_imagen(fila["Imagen URL"], ruta)
                nuevas += 1
            except Exception:
                pendientes.append(fila)

    fallos = []
    if pendientes:
        print(f"{len(pendientes)} fotos con la URL del HTML caducada: buscándolas en Google Maps...")
        urls = asyncio.run(_urls_frescas([f["Nombre"] for f in pendientes]))
        for fila in pendientes:
            if not urls.get(fila["Nombre"]):
                fallos.append(f"{fila['Nombre']}: no encontrada en Google Maps")
                continue
            try:
                bajar_imagen(urls[fila["Nombre"]], IMAGENES_DIR / nombre_archivo_imagen(fila["Nombre"]))
                nuevas += 1
            except Exception as e:
                fallos.append(f"{fila['Nombre']}: {str(e).split(' for url')[0]}")

    for fila in filas:
        archivo = nombre_archivo_imagen(fila["Nombre"])
        fila["Imagen archivo"] = archivo if (IMAGENES_DIR / archivo).exists() else ""

    print(f"Imágenes: {nuevas} descargadas, {sum(1 for f in filas if f['Imagen archivo'])}/{len(filas)} con foto.")
    if fallos:
        print(f"  - [AVISO] {len(fallos)} no se han podido descargar:")
        for f in fallos:
            print(f"      · {f}")


def extraer_restaurantes(soup):
    buttons = soup.find_all("button", class_="SMP2wb")
    print(f"Encontrados {len(buttons)} restaurantes en el HTML.")

    filas = []
    for b in buttons:
        # --- Nombre ---
        name_div = b.find("div", class_="fontHeadlineSmall")
        nombre = clean(name_div.get_text()) if name_div else ""

        # --- Puntuación y nº de reseñas (desde el aria-label) ---
        rating_span = b.find("span", class_="ZkP5Je")
        puntuacion, resenas = None, None
        if rating_span and rating_span.get("aria-label"):
            m = re.match(
                r"([\d,]+)\s+estrellas\s+([\d.,]+)\s+rese",
                rating_span["aria-label"],
            )
            if m:
                puntuacion = float(m.group(1).replace(",", "."))
                resenas = int(m.group(2).replace(".", ""))

        # --- Imagen (miniatura del restaurante) ---
        img_tag = b.find("img", class_="WkIe8")
        imagen_url = img_tag.get("src", "") if img_tag else ""

        # --- Segundo bloque IIrLbb: precio + cocina, o "Cerrado..." ---
        iirlbbs = b.find_all("div", class_="IIrLbb")
        precio, cocina, estado = "", "", ""

        if len(iirlbbs) >= 2:
            block = iirlbbs[1]
            texto_bloque = clean(block.get_text())

            if "cerrado" in texto_bloque.lower():
                # Caso: "Cerrado temporalmente" / "Cerrado permanentemente"
                estado = texto_bloque
            else:
                price_span = block.find("span", attrs={"role": "img"})
                if price_span:
                    precio = clean(price_span.get_text())

                spans = block.find_all("span", recursive=False)
                for sp in spans:
                    if sp.get("role") != "img":
                        text = clean(sp.get_text()).replace("·", "").strip()
                        if text:
                            cocina = text

        filas.append(
            {
                "Nombre": nombre,
                "Puntuación": puntuacion,
                "Nº Reseñas": resenas,
                "Rango de precios": precio,
                "Tipo de cocina": cocina,
                "Estado": estado,
                "Imagen URL": imagen_url,
            }
        )

    return filas


def guardar_excel(filas, ruta_salida):
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Restaurantes"

    headers = list(filas[0].keys()) if filas else [
        "Nombre", "Puntuación", "Nº Reseñas", "Rango de precios", "Tipo de cocina", "Estado", "Imagen URL",
        "Imagen archivo",
    ]
    ws.append(headers)
    for fila in filas:
        ws.append([fila[h] for h in headers])

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")

    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions

    anchos = {"A": 45, "B": 12, "C": 12, "D": 18, "E": 20, "F": 25, "G": 50, "H": 40}
    for col, w in anchos.items():
        ws.column_dimensions[col].width = w

    wb.save(ruta_salida)
    print(f"\nGuardado: {ruta_salida}")
    print(f"  - {len(filas)} restaurantes")

    cerrados = [f for f in filas if f["Estado"]]
    if cerrados:
        print(f"  - [AVISO] {len(cerrados)} restaurantes aparecen como cerrados:")
        for f in cerrados:
            print(f"      · {f['Nombre']} — {f['Estado']}")


async def _abrir_chrome_con_sesion(p):
    """Chrome con un perfil propio (PERFIL_CHROME_DIR) que recuerda la sesión de Google entre
    ejecuciones. Usa el Chrome instalado si lo hay (Google pone menos pegas al iniciar sesión
    en él que en el Chromium de Playwright) y, si no, el de Playwright.
    Ojo: la sesión NO se puede llevar a otro ordenador (p. ej. a GitHub Actions): Chrome la ata a su perfil y
    Google anula la que se usa desde un servidor de otro sitio (lo probamos)."""
    opciones = dict(user_data_dir=str(PERFIL_CHROME_DIR), headless=False, locale="es-ES",
                    viewport={"width": 1440, "height": 900},
                    args=["--disable-blink-features=AutomationControlled"],
                    # sin las marcas de "navegador automatizado", Google pone menos pegas al iniciar sesión
                    ignore_default_args=["--enable-automation", "--no-sandbox"])
    if SIN_VENTANA:
        # Chrome normal con la ventana fuera de la pantalla. NO el modo headless: sin ventana Google no da la
        # sesión, y disfrazarlo de Chrome normal hizo que Google la anulara del todo (lo probamos).
        opciones["args"] = opciones["args"] + ["--window-position=-32000,-32000"]
    try:
        return await p.chromium.launch_persistent_context(channel="chrome", **opciones)
    except Exception:
        return await p.chromium.launch_persistent_context(**opciones)


async def _con_sesion(page):
    return await page.locator('a[href*="accounts.google.com/SignOutOptions"]').count() > 0


async def _asegurar_sesion(page, url):
    """Con sesión, la lista trae precio y demás datos de cada restaurante; sin ella, no.
    Si no hay sesión guardada, abre el inicio de sesión de Google y espera a que entres."""
    try:   # el botón de la cuenta tarda unos segundos en aparecer tras cargar Maps
        await page.locator('a[href*="accounts.google.com/SignOutOptions"]').first.wait_for(state="attached",
                                                                                          timeout=15000)
        return
    except Exception:
        pass
    if SIN_VENTANA:   # no hay ventana donde iniciar sesión
        raise SystemExit("ERROR: no hay sesión de Google en el Chrome del procesador. Ejecuta una vez "
                         "'python obtener_datos/procesar_solicitudes.py' (sin --sin-ventana) e inicia sesión.")
    print("\n  No hay sesión de Google guardada. Inicia sesión en la ventana de Chrome que se ha abierto")
    print(f"  (solo hace falta la primera vez). Esperando hasta {ESPERA_INICIO_SESION_S // 60} minutos...\n")
    # Google solo deja volver a sus propios dominios tras iniciar sesión (con el enlace corto
    # maps.app.goo.gl da error 400): se vuelve a Maps y luego se abre la lista.
    await page.goto("https://accounts.google.com/ServiceLogin?hl=es&continue="
                    + urllib.parse.quote("https://www.google.com/maps?hl=es", safe=""))
    for _ in range(ESPERA_INICIO_SESION_S // 2):
        await page.wait_for_timeout(2000)
        if "google.com/maps" in page.url and await _con_sesion(page):
            print("  Sesión iniciada; queda guardada para las próximas veces.")
            await page.goto(url, wait_until="load")
            await page.wait_for_timeout(2500)
            return
    raise SystemExit("ERROR: no se ha iniciado sesión en Google a tiempo. Vuelve a ejecutar el script.")


async def _descargar_html_lista(url):
    """Abre tu lista de Google Maps con tu sesión, hace scroll hasta cargar todos los restaurantes,
    los recorre uno a uno para que Google rellene sus datos y devuelve el HTML de la página."""
    from playwright.async_api import async_playwright
    from extraer_resenas import aceptar_cookies

    async with async_playwright() as p:
        contexto = await _abrir_chrome_con_sesion(p)
        page = contexto.pages[0] if contexto.pages else await contexto.new_page()
        await page.goto(url, wait_until="load")
        await page.wait_for_timeout(2500)   # Google puede redirigir a la pantalla de cookies después de cargar
        await aceptar_cookies(page)
        await page.wait_for_load_state("load")
        await _asegurar_sesion(page, url)

        entradas = page.locator(SELECTOR_ENTRADA_LISTA)
        try:
            await entradas.first.wait_for(timeout=30000)
        except Exception:
            html = await page.content()   # sin restaurantes: main() lo guarda aparte y avisa
            await contexto.close()
            return html

        # 1) Cargar todas las entradas (la lista llega de 20 en 20 al llegar abajo).
        cargados, rondas_sin_nuevos = 0, 0
        for _ in range(MAX_RONDAS_SCROLL):
            await entradas.last.scroll_into_view_if_needed()
            await page.wait_for_timeout(1500)
            actuales = await entradas.count()
            rondas_sin_nuevos = rondas_sin_nuevos + 1 if actuales == cargados else 0
            cargados = actuales
            print(f"  {cargados} restaurantes cargados...")
            if rondas_sin_nuevos >= RONDAS_SIN_NUEVOS:
                break

        # 2) Recorrerlas una a una: Google solo rellena los datos de las entradas que pasan por pantalla.
        print("  Recorriendo la lista para que se carguen los datos de cada restaurante...")
        for i in range(cargados):
            await entradas.nth(i).scroll_into_view_if_needed()
            await page.wait_for_timeout(PAUSA_POR_ENTRADA_MS)
        incompletas = [i for i in range(cargados) if not await entradas.nth(i).locator("span.ZkP5Je").count()]
        for i in incompletas:   # segunda oportunidad, más despacio
            await entradas.nth(i).scroll_into_view_if_needed()
            try:
                await entradas.nth(i).locator("span.ZkP5Je").wait_for(timeout=2000)
            except Exception:
                pass

        html = await page.content()
        await contexto.close()
    return html


def main():
    import argparse
    parser = argparse.ArgumentParser(description="Lista de Google Maps -> restaurantes_v1.xlsx")
    parser.add_argument("--url", help="enlace de tu lista compartida de Google Maps (Compartir -> copiar "
                                      "enlace). Si se da, descarga la lista sola; si no, lee origen/restaurantes.html")
    args = parser.parse_args()

    if args.url:
        print("Abriendo la lista en Google Maps (se abre una ventana de Chrome; no la cierres)...")
        html = asyncio.run(_descargar_html_lista(args.url))
        if not extraer_restaurantes(BeautifulSoup(html, "lxml")):
            # No se pisa el HTML bueno que hubiera: lo descargado se guarda aparte para revisarlo.
            HTML_DESCARGA_FALLIDA.write_text(html, encoding="utf-8")
            raise SystemExit(
                f"ERROR: la página descargada no tiene ningún restaurante. Comprueba que el enlace es de "
                f"una lista compartida (Compartir -> copiar enlace) y que se abre sin iniciar sesión. "
                f"Lo descargado está en '{HTML_DESCARGA_FALLIDA}'; '{INPUT_HTML.name}' no se ha tocado."
            )
        INPUT_HTML.parent.mkdir(parents=True, exist_ok=True)
        INPUT_HTML.write_text(html, encoding="utf-8")
        print(f"Lista guardada en {INPUT_HTML}")

    if not INPUT_HTML.exists():
        raise SystemExit(
            f"ERROR: no encuentro '{INPUT_HTML}'. Usa --url con el enlace de tu lista compartida, o "
            f"guarda el HTML de tu lista de Google Maps en esa ruta."
        )

    with open(INPUT_HTML, encoding="utf-8") as f:
        soup = BeautifulSoup(f, "lxml")

    filas = extraer_restaurantes(soup)
    resumen = {campo: sum(1 for f in filas if f[campo] not in ("", None))
               for campo in ("Puntuación", "Rango de precios", "Tipo de cocina", "Imagen URL")}
    print("Con datos: " + " | ".join(f"{campo} {n}/{len(filas)}" for campo, n in resumen.items()))
    descargar_imagenes(filas)
    guardar_excel(filas, OUTPUT_XLSX)


if __name__ == "__main__":
    main()