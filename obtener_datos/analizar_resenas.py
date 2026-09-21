"""
Analiza el sentimiento de las reseñas (buena/mala) y detecta menciones de
platos concretos valorados positiva o negativamente.

Método 100% local y gratuito:
  - Sentimiento general: modelo en español de la librería `pysentimiento`.
  - Detección de platos: diccionario editable + fragmentación de frases.

LIMITACIÓN: solo se detectan platos cuyo nombre (o similar) esté en
DICCIONARIO_PLATOS más abajo. Añade términos libremente si ves que se
escapan platos de tu lista de restaurantes.

USO:
    1. Instala dependencias (la primera vez tardará un poco, instala
       transformers/torch):
         pip install pysentimiento pandas openpyxl

    2. Ejecuta (usa como entrada el excel generado por extraer_resenas.py):
         python analizar_resenas.py

    La primera ejecución descargará el modelo de sentimiento (~500MB) de
    internet, así que necesitas conexión la primera vez. Ejecuciones
    posteriores usan el modelo ya descargado (caché local).
"""

import re
import sys
import time
import unicodedata
from collections import defaultdict

import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill


# ----------------------- CONFIGURACIÓN ----------------------- #

INPUT_XLSX = "restaurantes_con_resenas.xlsx"     # salida del script anterior
INPUT_SHEET = "Reseñas"
OUTPUT_XLSX = "analisis_resenas.xlsx"

LIMITE_RESENAS = None   # pon un número (ej. 50) para probar rápido con pocas reseñas antes de lanzar todo

# --------------------------------------------------------------- #


# Diccionario de platos/ingredientes en español. AÑADE AQUÍ lo que veas
# que falta para tu lista concreta de restaurantes.
DICCIONARIO_PLATOS = [
    # Genéricos
    "pasta", "pizza", "hamburguesa", "burger", "ensalada", "sopa", "postre",
    "tarta", "helado", "pan", "arroz", "carne", "pescado", "marisco",
    "pollo", "cordero", "cerdo", "ternera", "verduras", "patatas",
    "patatas fritas", "croquetas", "tortilla", "gazpacho", "paella",
    "sándwich", "bocadillo", "empanada", "focaccia", "alitas", "costillas",
    "solomillo", "chuletón", "buey", "entrecot", "salmón", "atún", "pulpo",
    "gambas", "langostinos", "vieiras", "mejillones", "almejas", "ostras",
    "bacalao",
    # Japonesa
    "sushi", "sashimi", "ramen", "tempura", "gyoza", "udon", "yakitori",
    "poke", "wakame", "edamame", "miso", "teriyaki", "katsu", "nigiri",
    "maki", "temaki",
    # Tailandesa
    "pad thai", "curry", "tom yum", "satay", "wok",
    # Italiana
    "risotto", "lasaña", "carbonara", "burrata", "tiramisú",
    "prosciutto", "gnocchi", "ravioli", "spaghetti", "espaguetis",
    # Coreana
    "bibimbap", "kimchi", "bulgogi", "japchae",
    # Libanesa / Griega
    "hummus", "falafel", "tabulé", "shawarma", "gyros", "moussaka",
    "tzatziki", "pita",
    # Mexicana / Tex-Mex
    "tacos", "guacamole", "nachos", "burrito", "quesadilla", "fajitas",
]

# Palabras/frases que usamos para partir una reseña en fragmentos más
# pequeños, de forma que si hay opiniones contrapuestas sobre distintos
# platos no se mezclen en un mismo fragmento.
SEPARADORES = r"[.!?;\n]|,?\s*\b(?:pero|aunque|sin embargo|mientras que)\b"


def cargar_resenas(ruta, hoja):
    df = pd.read_excel(ruta, sheet_name=hoja)
    columnas_necesarias = {"Restaurante", "Texto"}
    if not columnas_necesarias.issubset(df.columns):
        raise ValueError(f"Faltan columnas {columnas_necesarias} en la hoja '{hoja}'")
    df = df.dropna(subset=["Texto"])
    df["Texto"] = df["Texto"].astype(str)
    if LIMITE_RESENAS:
        df = df.head(LIMITE_RESENAS)
    return df


def crear_analizador():
    print("Cargando modelo de sentimiento en español (pysentimiento)...")
    print("(la primera vez puede tardar varios minutos por la descarga)")
    from pysentimiento import create_analyzer
    analyzer = create_analyzer(task="sentiment", lang="es")
    return analyzer


def mapear_sentimiento(resultado):
    """Convierte la salida de pysentimiento (POS/NEG/NEU) a texto legible."""
    etiqueta = resultado.output
    return {"POS": "Buena", "NEG": "Mala", "NEU": "Neutra"}.get(etiqueta, "Neutra")


def dividir_en_fragmentos(texto):
    fragmentos = re.split(SEPARADORES, texto, flags=re.IGNORECASE)
    return [f.strip() for f in fragmentos if f and f.strip()]


def normalizar(texto):
    """Quita acentos para comparar de forma más permisiva."""
    nfkd = unicodedata.normalize("NFKD", texto.lower())
    return "".join(c for c in nfkd if not unicodedata.combining(c))


# Diccionario normalizado (sin acentos) -> término original, para buscar
DICCIONARIO_NORM = {normalizar(plato): plato for plato in DICCIONARIO_PLATOS}


def buscar_platos_en_fragmento(fragmento):
    frag_norm = normalizar(fragmento)
    encontrados = []
    for plato_norm, plato_original in DICCIONARIO_NORM.items():
        patron = r"\b" + re.escape(plato_norm) + r"\b"
        if re.search(patron, frag_norm):
            encontrados.append(plato_original)
    return encontrados


def analizar_resenas(df, analyzer):
    filas_resenas = []
    filas_platos = []

    total = len(df)
    t0 = time.time()

    for idx, fila in enumerate(df.itertuples(index=False), start=1):
        texto = fila.Texto
        restaurante = fila.Restaurante
        autor = getattr(fila, "Autor", "")

        # --- Sentimiento general de la reseña completa ---
        resultado_general = analyzer.predict(texto)
        sentimiento_general = mapear_sentimiento(resultado_general)

        filas_resenas.append({
            "Restaurante": restaurante,
            "Autor": autor,
            "Texto": texto,
            "Sentimiento reseña": sentimiento_general,
            "Confianza": round(max(resultado_general.probas.values()), 2),
        })

        # --- Búsqueda de platos por fragmento ---
        for fragmento in dividir_en_fragmentos(texto):
            platos = buscar_platos_en_fragmento(fragmento)
            if not platos:
                continue

            resultado_frag = analyzer.predict(fragmento)
            sentimiento_frag = mapear_sentimiento(resultado_frag)

            for plato in platos:
                filas_platos.append({
                    "Restaurante": restaurante,
                    "Plato": plato,
                    "Sentimiento": sentimiento_frag,
                    "Fragmento": fragmento,
                    "Autor reseña": autor,
                })

        if idx % 25 == 0 or idx == total:
            transcurrido = time.time() - t0
            print(f"  Procesadas {idx}/{total} reseñas ({transcurrido:.0f}s)...")

    return filas_resenas, filas_platos


def construir_resumen(filas_resenas, filas_platos):
    resumen = defaultdict(lambda: {
        "buenas": 0, "malas": 0, "neutras": 0,
        "platos_buenos": defaultdict(int), "platos_malos": defaultdict(int),
    })

    for fila in filas_resenas:
        r = resumen[fila["Restaurante"]]
        if fila["Sentimiento reseña"] == "Buena":
            r["buenas"] += 1
        elif fila["Sentimiento reseña"] == "Mala":
            r["malas"] += 1
        else:
            r["neutras"] += 1

    for fila in filas_platos:
        r = resumen[fila["Restaurante"]]
        if fila["Sentimiento"] == "Buena":
            r["platos_buenos"][fila["Plato"]] += 1
        elif fila["Sentimiento"] == "Mala":
            r["platos_malos"][fila["Plato"]] += 1

    filas_resumen = []
    for restaurante, datos in resumen.items():
        top_buenos = sorted(datos["platos_buenos"].items(), key=lambda x: -x[1])[:5]
        top_malos = sorted(datos["platos_malos"].items(), key=lambda x: -x[1])[:5]
        filas_resumen.append({
            "Restaurante": restaurante,
            "Reseñas buenas": datos["buenas"],
            "Reseñas malas": datos["malas"],
            "Reseñas neutras": datos["neutras"],
            "Platos mejor valorados": ", ".join(f"{p} ({n})" for p, n in top_buenos),
            "Platos peor valorados": ", ".join(f"{p} ({n})" for p, n in top_malos),
        })

    filas_resumen.sort(key=lambda f: f["Restaurante"])
    return filas_resumen


def guardar_excel(filas_resenas, filas_platos, filas_resumen, ruta_salida):
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "Reseñas analizadas"
    if filas_resenas:
        headers = list(filas_resenas[0].keys())
        ws1.append(headers)
        for fila in filas_resenas:
            ws1.append([fila[h] for h in headers])

    ws2 = wb.create_sheet("Platos mencionados")
    if filas_platos:
        headers = list(filas_platos[0].keys())
        ws2.append(headers)
        for fila in filas_platos:
            ws2.append([fila[h] for h in headers])

    ws3 = wb.create_sheet("Resumen por restaurante")
    if filas_resumen:
        headers = list(filas_resumen[0].keys())
        ws3.append(headers)
        for fila in filas_resumen:
            ws3.append([fila[h] for h in headers])

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial")
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions

    anchos = {
        ws1: {"A": 30, "B": 20, "C": 60, "D": 16, "E": 12},
        ws2: {"A": 30, "B": 18, "C": 14, "D": 60, "E": 20},
        ws3: {"A": 30, "B": 14, "C": 14, "D": 14, "E": 45, "F": 45},
    }
    for ws, widths in anchos.items():
        for col, w in widths.items():
            ws.column_dimensions[col].width = w

    wb.save(ruta_salida)
    print(f"\nGuardado: {ruta_salida}")
    print(f"  - {len(filas_resenas)} reseñas analizadas")
    print(f"  - {len(filas_platos)} menciones de platos detectadas")
    print(f"  - {len(filas_resumen)} restaurantes resumidos")


def main():
    df = cargar_resenas(INPUT_XLSX, INPUT_SHEET)
    print(f"Leídas {len(df)} reseñas de '{INPUT_XLSX}' (hoja '{INPUT_SHEET}').")

    analyzer = crear_analizador()

    filas_resenas, filas_platos = analizar_resenas(df, analyzer)
    filas_resumen = construir_resumen(filas_resenas, filas_platos)

    guardar_excel(filas_resenas, filas_platos, filas_resumen, OUTPUT_XLSX)


if __name__ == "__main__":
    main()
